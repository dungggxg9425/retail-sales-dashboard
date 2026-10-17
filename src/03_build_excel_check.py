"""
Bước 3 – Tạo workbook Excel đối chiếu số liệu: dữ liệu đã làm sạch + bảng tổng hợp bằng công thức SUMIFS.
Dùng để kiểm tra kết quả các measure DAX trên dashboard (số trên dashboard phải khớp với sheet Doi_chieu).
Đầu ra: excel/doi_chieu_so_lieu.xlsx
Chạy  : python src/03_build_excel_check.py
"""
from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
M = ROOT / "data" / "model"
OUT = ROOT / "excel"
OUT.mkdir(exist_ok=True)

# Dữ liệu phẳng: bảng sự kiện + thuộc tính cần lọc
f = (pd.read_csv(M / "fact_sales.csv", parse_dates=["order_date"])
       .merge(pd.read_csv(M / "dim_product.csv"), on="product_key")
       .merge(pd.read_csv(M / "dim_geography.csv"), on="geo_key")
       .merge(pd.read_csv(M / "dim_customer.csv"), on="customer_id"))
f["year"] = f["order_date"].dt.year
cols = ["row_id", "order_id", "order_date", "year", "segment", "region", "state", "category", "sub_category",
        "product_name", "sales", "quantity", "discount", "profit"]
f = f[cols].sort_values("row_id")

FONT = "Arial"
H = Font(name=FONT, bold=True, color="FFFFFF")
HFILL = PatternFill("solid", fgColor="1F1F1F")
B = Font(name=FONT, bold=True)
N = Font(name=FONT)
thin = Border(bottom=Side(style="thin", color="BFBFBF"))

wb = Workbook()

# ---------------- Sheet Data
ws = wb.active
ws.title = "Data"
ws.append(cols)
for r in f.itertuples(index=False):
    ws.append([r.row_id, r.order_id, r.order_date.to_pydatetime(), r.year, r.segment, r.region, r.state,
               r.category, r.sub_category, r.product_name, r.sales, r.quantity, r.discount, r.profit])
for c in ws[1]:
    c.font, c.fill = H, HFILL
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.font = N
    row[2].number_format = "yyyy-mm-dd"
    row[10].number_format = row[13].number_format = "#,##0.00"
    row[12].number_format = "0%"
widths = [8, 16, 12, 7, 12, 9, 16, 15, 13, 50, 11, 9, 9, 11]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions
last = ws.max_row
rng = lambda col: f"Data!${col}$2:${col}${last}"   # vùng dữ liệu cố định
SALES, PROFIT, YEAR, CAT, REG, DISC, ORDER = rng("K"), rng("N"), rng("D"), rng("H"), rng("F"), rng("M"), rng("B")

# ---------------- Sheet Doi_chieu
wc = wb.create_sheet("Doi_chieu")
wc["A1"] = "Đối chiếu số liệu dashboard – Superstore (tính bằng SUMIFS trên sheet Data)"
wc["A1"].font = Font(name=FONT, bold=True, size=13)
wc["A2"] = "Mọi ô số là công thức; sửa dữ liệu ở sheet Data thì bảng tự cập nhật. Biên LN = Lợi nhuận / Doanh thu."
wc["A2"].font = Font(name=FONT, italic=True, color="595959")


def header(row, labels):
    for j, lab in enumerate(labels):
        c = wc.cell(row=row, column=1 + j, value=lab)
        c.font, c.fill = H, HFILL
        c.alignment = Alignment(horizontal="center")


def fmt(row, money_cols=(2, 3), pct_cols=(4,)):
    for j in money_cols:
        wc.cell(row=row, column=j).number_format = "#,##0.00"
    for j in pct_cols:
        wc.cell(row=row, column=j).number_format = "0.0%"
    for j in range(1, 6):
        wc.cell(row=row, column=j).font = N
        wc.cell(row=row, column=j).border = thin


r = 4
# Tổng
wc.cell(row=r, column=1, value="Tổng toàn kỳ").font = B
r += 1
header(r, ["Chỉ số", "Doanh thu", "Lợi nhuận", "Biên LN", "Số dòng"])
r += 1
wc.cell(row=r, column=1, value="2014–2017")
wc.cell(row=r, column=2, value=f"=SUM({SALES})")
wc.cell(row=r, column=3, value=f"=SUM({PROFIT})")
wc.cell(row=r, column=4, value=f"=C{r}/B{r}")
wc.cell(row=r, column=5, value=f"=COUNTA({ORDER})")
fmt(r)
r += 2

# Theo năm + YoY
wc.cell(row=r, column=1, value="Theo năm").font = B
r += 1
header(r, ["Năm", "Doanh thu", "Lợi nhuận", "Biên LN", "Tăng trưởng DT YoY"])
r += 1
first = r
for yr in sorted(f["year"].unique()):
    wc.cell(row=r, column=1, value=int(yr))
    wc.cell(row=r, column=2, value=f"=SUMIFS({SALES},{YEAR},A{r})")
    wc.cell(row=r, column=3, value=f"=SUMIFS({PROFIT},{YEAR},A{r})")
    wc.cell(row=r, column=4, value=f"=C{r}/B{r}")
    wc.cell(row=r, column=5, value="" if r == first else f"=B{r}/B{r-1}-1")
    fmt(r, pct_cols=(4, 5))
    r += 1
r += 1


def by_dim(title, label, values, crit_range):
    global r
    wc.cell(row=r, column=1, value=title).font = B
    r += 1
    header(r, [label, "Doanh thu", "Lợi nhuận", "Biên LN", "Tỷ trọng DT"])
    r += 1
    start = r
    for v in values:
        wc.cell(row=r, column=1, value=v)
        wc.cell(row=r, column=2, value=f"=SUMIFS({SALES},{crit_range},A{r})")
        wc.cell(row=r, column=3, value=f"=SUMIFS({PROFIT},{crit_range},A{r})")
        wc.cell(row=r, column=4, value=f"=C{r}/B{r}")
        wc.cell(row=r, column=5, value=f"=B{r}/SUM($B${start}:$B${start + len(values) - 1})")
        fmt(r, pct_cols=(4, 5))
        r += 1
    r += 1


by_dim("Theo danh mục", "Danh mục", sorted(f["category"].unique()), CAT)
by_dim("Theo khu vực", "Khu vực", sorted(f["region"].unique()), REG)

# Theo mức chiết khấu (điều kiện khoảng bằng SUMIFS nhiều điều kiện)
wc.cell(row=r, column=1, value="Theo mức chiết khấu").font = B
r += 1
header(r, ["Mức chiết khấu", "Doanh thu", "Lợi nhuận", "Biên LN", "Số dòng"])
r += 1
bands = [("0%", ">=0", "<=0"), ("1–20%", ">0", "<=0.2"), ("21–40%", ">0.2", "<=0.4"), ("> 40%", ">0.4", "<=1")]
for lab, lo, hi in bands:
    wc.cell(row=r, column=1, value=lab)
    wc.cell(row=r, column=2, value=f'=SUMIFS({SALES},{DISC},"{lo}",{DISC},"{hi}")')
    wc.cell(row=r, column=3, value=f'=SUMIFS({PROFIT},{DISC},"{lo}",{DISC},"{hi}")')
    wc.cell(row=r, column=4, value=f"=C{r}/B{r}")
    wc.cell(row=r, column=5, value=f'=COUNTIFS({DISC},"{lo}",{DISC},"{hi}")')
    fmt(r)
    wc.cell(row=r, column=5).number_format = "#,##0"
    r += 1

for col, w in zip("ABCDE", [22, 16, 16, 11, 20]):
    wc.column_dimensions[col].width = w
wb.move_sheet("Doi_chieu", offset=-1)
wb.active = 0
wb.save(OUT / "doi_chieu_so_lieu.xlsx")
print("Đã tạo", OUT / "doi_chieu_so_lieu.xlsx")
